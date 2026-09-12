#!/usr/bin/env python3
"""
Template 5: Executive KPI Scorecard & Strategy Matrix
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
4 Strategic Performance Pillars, KPI metric cards with variance trends (+/- %),
RAG health status pills, diagnostic deep-dive panels, and executive action grid.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from ..generator_core import (
    PALETTES, LW, bx, fl, ci, tx, tp, shadow, dash_bd,
    conn_arrow, conn_line, ad, ar, ln, emote_r, AnimationEngine
)

def build_kpi_scorecard(data, static_dir, out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.50)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    eng = AnimationEngine()
    P1, P2, P3, P4 = [], [], [], []

    branding = data.get("branding", {})
    palette_key = branding.get("palette", "navy_corporate")
    C = PALETTES.get(palette_key, PALETTES["navy_corporate"])

    def get_emote(name):
        p = os.path.join(static_dir, "emotes", f"{name}.gif")
        if not os.path.exists(p):
            p = os.path.join(static_dir, "emotes", "tb_dash.gif")
        return p

    logo_file = branding.get("logo_file", "philips.png")
    logo_path = os.path.join(static_dir, "logos", logo_file)

    canvas = fl(slide, 0, 0, 13.33, 7.50, C["canvas_bg"])

    # ─── 1. HEADER (Y: 0.00 to 0.62) ──────────────────────────────────
    hdr = fl(slide, 0, 0, 13.33, 0.62, C["hdr_bg"]); P1.append((hdr.shape_id, 0, 'fade'))
    stripe = fl(slide, 0, 0.62, 13.33, 0.025, C["stripe"]); P1.append((stripe.shape_id, 150, 'wr'))

    title_text = branding.get("title", "EXECUTIVE KPI SCORECARD & PERFORMANCE DASHBOARD")
    sub_text = branding.get("subtitle", f"{branding.get('company_name', 'GLOBAL OPERATIONS')}  •  CFO QUARTERLY PERFORMANCE REVIEW & SLA HEALTH MONITOR")

    ht = tx(slide, 0.40, 0.07, 9.00, 0.50)
    tp(ht.text_frame, title_text, sz=18, b=True, c=C["card_bg"])
    tp(ht.text_frame, sub_text, sz=8.5, c=C["t4"], sb=2)
    P1.append((ht.shape_id, 250, 'fade'))

    pbadge = bx(slide, 11.35, 0.10, 1.60, 0.42, C["card_bg"], None, r=4000)
    shadow(pbadge, 8000, 4000, 7000)
    if os.path.exists(logo_path):
        plogo = slide.shapes.add_picture(logo_path, Inches(11.45), Inches(0.15), Inches(1.40), Inches(0.32))
        P1 += [(pbadge.shape_id, 350, 'fade'), (plogo.shape_id, 400, 'fade')]
    else:
        tp(pbadge.text_frame, branding.get("company_name", "OPERATIONS"), sz=12, b=True, c=C["accent"], a=PP_ALIGN.CENTER)
        P1.append((pbadge.shape_id, 350, 'fade'))

    # ─── 2. 4 TOP KPI SUMMARY CARDS (Y: 0.74 to 2.20) ──────────────────
    kpis = data.get("top_kpis", [
        {"title": "Straight-Through Processing", "val": "84.2%", "target": "Target: >80%", "trend": "+12.4% MoM", "status": "GREEN", "emote": "tb_write"},
        {"title": "Unmatched Open Delta", "val": "$4.1M", "target": "Target: <$5.0M", "trend": "–38.1% YoY", "status": "GREEN", "emote": "cash"},
        {"title": "Average Resolution TAT", "val": "46h", "target": "SLA: <72h", "trend": "–26h vs H1", "status": "GREEN", "emote": "waiting"},
        {"title": "Escalation Rate (>96h)", "val": "2.8%", "target": "Target: <3.0%", "trend": "–1.4% MoM", "status": "GREEN", "emote": "esc"}
    ])

    card_w = 2.92
    card_gap = 0.28
    card_xs = [0.40, 0.40 + (card_w + card_gap), 0.40 + (card_w + card_gap) * 2, 0.40 + (card_w + card_gap) * 3]

    for k_idx, kpi in enumerate(kpis[:4]):
        kx = card_xs[k_idx]
        kb = bx(slide, kx, 0.74, card_w, 1.46, C["card_bg"], C["card_bd"], r=5000)
        shadow(kb, 10000, 5000, 7000)

        # RAG Status Indicator Bar
        rag_c = C["stripe"] if kpi.get("status") == "GREEN" else (C["amber"] if kpi.get("status") == "AMBER" else C["red"])
        kbar = fl(slide, kx, 0.82, 0.045, 1.30, rag_c)

        eg, ee = emote_r(slide, get_emote(kpi.get("emote", "tb_dash")), kx, 0.74, card_w, 1.46, C["accent_g"], 0.48)

        ktx = tx(slide, kx + 0.14, 0.82, card_w - 0.72, 1.30)
        tp(ktx.text_frame, kpi.get("title", "Metric"), sz=8.6, b=True, c=C["t2"])
        p_v = tp(ktx.text_frame, kpi.get("val", "0.0%"), sz=18, b=True, c=C["t1"], sb=2.0)
        p_tr = tp(ktx.text_frame, f"{kpi.get('trend', '')}  |  {kpi.get('target', '')}", sz=7.0, b=True, c=rag_c, sb=2.0)

        d_k = 400 + k_idx * 120
        P1 += [(kb.shape_id, d_k, 'fade'), (kbar.shape_id, d_k + 10, 'fade'),
               (eg.shape_id, d_k + 20, 'fade'), (ee.shape_id, d_k + 30, 'ep'), (ktx.shape_id, d_k + 40, 'fade')]

    # ─── 3. STRATEGIC PILLAR DEEP DIVE (Y: 2.36 to 5.40, H: 3.04) ─────
    pillars = data.get("pillars", [
        {
            "name": "1. Operational Excellence",
            "score": "94 / 100",
            "items": [
                ("ERP Auto-Triage Accuracy", "91.8%", C["stripe"]),
                ("HWI Hand-off Latency", "1.2h", C["stripe"]),
                ("Invoice OCR Match Rate", "88.4%", C["amber"])
            ],
            "emote": "review"
        },
        {
            "name": "2. Financial Risk & Controls",
            "score": "98 / 100",
            "items": [
                ("SOX Audit Trail Completeness", "100%", C["stripe"]),
                ("AP-AR Asymmetry Zero-Days", "99.4%", C["stripe"]),
                ("Unapplied Cash Write-offs", "$0.00", C["stripe"])
            ],
            "emote": "tb_audit"
        },
        {
            "name": "3. Counterparty SLA Health",
            "score": "89 / 100",
            "items": [
                ("Bilateral First-Pass Yield", "82.6%", C["amber"]),
                ("Email Action Response Rate", "94.2%", C["stripe"]),
                ("Level-3 Escalation Drops", "0.2%", C["stripe"])
            ],
            "emote": "notif"
        },
        {
            "name": "4. Intelligent Automation",
            "score": "92 / 100",
            "items": [
                ("Self-Healing Journal Postings", "4,210/mo", C["stripe"]),
                ("RPA Extract Job Reliability", "99.9%", C["stripe"]),
                ("FTE Capacity Reinvested", "320 hrs", C["stripe"])
            ],
            "emote": "tb_robot"
        }
    ])

    pil_w = 2.92
    for p_idx, pil in enumerate(pillars[:4]):
        px = card_xs[p_idx]
        pb = bx(slide, px, 2.36, pil_w, 3.04, C["card_bg"], C["card_bd"], r=5000)
        shadow(pb, 10000, 5000, 7000)

        # Header Strip inside Pillar
        phb = fl(slide, px, 2.36, pil_w, 0.32, C["hdr_bg"])
        phtx = tx(slide, px + 0.12, 2.40, pil_w - 0.24, 0.24)
        tp(phtx.text_frame, pil["name"], sz=8.0, b=True, c=C["card_bg"])

        eg, ee = emote_r(slide, get_emote(pil.get("emote", "review")), px, 2.76, pil_w, 0.60, C["accent_g"], 0.44)

        # Score Pill
        sc_pill = bx(slide, px + 0.12, 2.80, 1.20, 0.26, C["accent_l"], C["card_bd"], r=3000)
        tp(sc_pill.text_frame, f"HEALTH: {pil.get('score', '95/100')}", sz=7.0, b=True, c=C["accent"], a=PP_ALIGN.CENTER)

        # Breakdown List
        itx = tx(slide, px + 0.12, 3.16, pil_w - 0.24, 2.10)
        for label, val, c_val in pil.get("items", []):
            p = itx.text_frame.add_paragraph()
            p.space_before = Pt(6.0)
            r1 = p.add_run(); r1.text = label + "\n"; r1.font.size = Pt(7.2); r1.font.color.rgb = C["t3"]; r1.font.name = 'Aptos'
            r2 = p.add_run(); r2.text = val; r2.font.size = Pt(9.6); r2.font.bold = True; r2.font.color.rgb = c_val; r2.font.name = 'Aptos'

        d_p = 600 + p_idx * 120
        P2 += [(pb.shape_id, d_p, 'fade'), (phb.shape_id, d_p + 10, 'fade'), (phtx.shape_id, d_p + 20, 'fade'),
               (sc_pill.shape_id, d_p + 30, 'zm'), (eg.shape_id, d_p + 40, 'fade'), (ee.shape_id, d_p + 50, 'ep'),
               (itx.shape_id, d_p + 60, 'fade')]

    # ─── 4. EXECUTIVE ACTION PLAN (Bottom: Y: 5.56 to 7.08, H: 1.52) ──
    act_box = bx(slide, 0.40, 5.56, 12.53, 1.52, C["card_bg"], C["card_bd"], r=5000)
    shadow(act_box, 12000, 6000, 8000)
    dash_bd(act_box, C["ctn_bd"], 1.5, 'lgDash')

    act_pill = bx(slide, 0.54, 5.46, 2.60, 0.22, C["hdr_bg"], None, r=3000)
    tp(act_pill.text_frame, "Q3 STRATEGIC INTERVENTIONS & ACCOUNTABILITY", sz=7.0, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    P3 += [(act_box.shape_id, 800, 'fade'), (act_pill.shape_id, 830, 'zm')]

    actions = data.get("actions", [
        ("EDI Direct Connector Rollout", "Deploy automated invoice drop adapter for top 15 paper entities.", "Owner: Integration Lead", "Due: Aug 15", C["blue"]),
        ("SLA Escalation Protocol Tuning", "Reduce L2 notification threshold from 96h to 72h across EMEA.", "Owner: Shared Services VP", "Due: Aug 30", C["amber"]),
        ("AI Auto-Clearing Rules Expansion", "Increase autonomous journal posting scope for items under $10k.", "Owner: Intelligent Automation CoE", "Due: Sep 15", C["teal"]),
        ("Bilateral Netting Policy Mandate", "Execute quarterly multilateral netting runs to minimize gross wire transfers.", "Owner: Treasury Director", "Due: Sep 30", C["rose"])
    ])

    for a_idx, (a_title, a_desc, a_own, a_due, a_col) in enumerate(actions[:4]):
        ax = card_xs[a_idx]
        ay = 5.76
        aw = card_w
        ah = 1.22

        ab = bx(slide, ax, ay, aw, ah, C["accent_l"], C["card_bd"], r=3000)
        atag = fl(slide, ax, ay + 0.06, 0.04, ah - 0.12, a_col)
        atx = tx(slide, ax + 0.10, ay + 0.06, aw - 0.16, ah - 0.12)
        tp(atx.text_frame, a_title, sz=8.4, b=True, c=C["t1"])
        tp(atx.text_frame, a_desc, sz=6.8, c=C["t3"], sb=1.5)
        tp(atx.text_frame, f"{a_own}  |  {a_due}", sz=6.6, b=True, c=a_col, sb=2.0)

        d_act = 860 + a_idx * 50
        P3 += [(ab.shape_id, d_act, 'fade'), (atag.shape_id, d_act + 10, 'fade'), (atx.shape_id, d_act + 20, 'fade')]

    eng.add(P1); eng.add(P2); eng.add(P3)
    eng.build(slide)
    prs.save(out_path)
    return out_path
