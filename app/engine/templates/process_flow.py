#!/usr/bin/env python3
"""
Template 1: Executive Process Flow & Decision Tree
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
ICA Archetype: 4-phase ribbon, left ingestion pipeline,
symmetrical decision tree tracks, convergence bus, and governance escalation.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from ..generator_core import (
    PALETTES, LW, bx, fl, ci, tx, tp, shadow, dash_bd,
    conn_arrow, conn_line, ad, ar, ln, emote_r, AnimationEngine
)

def build_process_flow(data, static_dir, out_path):
    """
    Builds the Process Flow & Decision Tree presentation.
    data: dict containing branding, palette, ribbon, ingestion, tracks, governance.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.50)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    eng = AnimationEngine()
    P1, P2, P3, P4, P5 = [], [], [], [], []

    branding = data.get("branding", {})
    palette_key = branding.get("palette", "navy_corporate")
    C = PALETTES.get(palette_key, PALETTES["navy_corporate"])

    # Emote path resolver
    def get_emote(name):
        p = os.path.join(static_dir, "emotes", f"{name}.gif")
        if not os.path.exists(p):
            p = os.path.join(static_dir, "emotes", "spec.gif")
        return p

    # Logo path resolver
    logo_file = branding.get("logo_file", "philips.png")
    logo_path = os.path.join(static_dir, "logos", logo_file)
    if not os.path.exists(logo_path):
        logo_path = os.path.join(static_dir, "logos", "philips.png")

    # Canvas background
    canvas = fl(slide, 0, 0, 13.33, 7.50, C["canvas_bg"])

    # ─── 1. HEADER (Y: 0.00 to 0.62) ──────────────────────────────────
    hdr = fl(slide, 0, 0, 13.33, 0.62, C["hdr_bg"]); P1.append((hdr.shape_id, 0, 'fade'))
    stripe = fl(slide, 0, 0.62, 13.33, 0.025, C["stripe"]); P1.append((stripe.shape_id, 150, 'wr'))

    title_text = branding.get("title", "ICA Reconciliation  |  AS-IS WORKFLOW BLUEPRINT")
    sub_text = branding.get("subtitle", f"{branding.get('company_name', 'PHILIPS')}  •  INTERCOMPANY ACCOUNTING PROCESS FLOW & GOVERNANCE")

    ht = tx(slide, 0.40, 0.07, 9.00, 0.50)
    tp(ht.text_frame, title_text, sz=18, b=True, c=C["card_bg"])
    tp(ht.text_frame, sub_text, sz=8.5, c=C["t4"], sb=2)
    P1.append((ht.shape_id, 250, 'fade'))

    # Logo Badge
    pbadge = bx(slide, 11.35, 0.10, 1.60, 0.42, C["card_bg"], None, r=4000)
    shadow(pbadge, 8000, 4000, 7000)
    if os.path.exists(logo_path):
        plogo = slide.shapes.add_picture(logo_path, Inches(11.45), Inches(0.15), Inches(1.40), Inches(0.32))
        P1 += [(pbadge.shape_id, 350, 'fade'), (plogo.shape_id, 400, 'fade')]
    else:
        tp(pbadge.text_frame, branding.get("company_name", "PHILIPS"), sz=12, b=True, c=C["accent"], a=PP_ALIGN.CENTER)
        P1.append((pbadge.shape_id, 350, 'fade'))

    # ─── 2. PIPELINE RIBBON (Y: 0.71 to 1.07) ─────────────────────────
    ribbon_items = data.get("ribbon", [
        {"num": "01", "name": "INGESTION & SCOPE", "sub": "Qlik Sense Extract & Filtering", "color": "blue"},
        {"num": "02", "name": "CLASSIFY & TRIAGE", "sub": "3-Track Operational Taxonomy", "color": "amber"},
        {"num": "03", "name": "RESOLVE BY CAUSE", "sub": "HWI, OCR, Forensic & Counterparty", "color": "teal"},
        {"num": "04", "name": "CLOSE THE LOOP", "sub": "Escalation Matrix & Governance", "color": "rose"}
    ])
    ribbon_widths = [2.70, 3.40, 2.70, 2.83]
    ribbon_xs = [0.40, 3.40, 7.10, 10.10]

    for i, item in enumerate(ribbon_items[:4]):
        x = ribbon_xs[i]
        w = ribbon_widths[i]
        ac = C.get(item.get("color", "blue"), C["accent"])
        s = bx(slide, x, 0.71, w, 0.36, C["card_bg"], C["card_bd"], r=4000); shadow(s, 10000, 5000, 7000)
        bd = bx(slide, x + 0.08, 0.75, 0.38, 0.28, ac, None, r=3000)
        bd.text_frame.margin_left = bd.text_frame.margin_right = 0
        tp(bd.text_frame, item.get("num", f"0{i+1}"), sz=9.0, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
        st = tx(slide, x + 0.52, 0.72, w - 0.56, 0.34)
        tp(st.text_frame, item.get("name", "STAGE"), sz=9.0, b=True, c=C["t1"])
        tp(st.text_frame, item.get("sub", ""), sz=6.8, c=C["t3"], sb=1)
        d = 550 + i * 150
        P1 += [(s.shape_id, d, 'fade'), (bd.shape_id, d + 30, 'zm'), (st.shape_id, d + 50, 'fade')]
        if i < 3:
            a = ar(slide, x + w + 0.06, 0.89, x + w + 0.24, C["t3"], w=LW); P1.append((a.shape_id, d + 90, 'wr'))

    # ─── 3. MAIN WORKFLOW CONTAINER (Y: 1.18 to 7.32, H: 6.14) ────────
    ctr = bx(slide, 0.40, 1.18, 12.53, 6.14, C["card_bg"], C["ctn_bd"], bw=1.8, r=10000)
    dash_bd(ctr, C["ctn_bd"], 1.8, 'lgDash'); shadow(ctr, 18000, 9000, 8000)
    P2.append((ctr.shape_id, 0, 'fade'))

    flow_badge = data.get("flow_badge", f"{branding.get('company_name', 'PHILIPS')} AS-IS FLOW")
    pill1 = bx(slide, 0.52, 1.08, 2.00, 0.24, C["amber"], None, r=5000)
    pill1.text_frame.margin_left = pill1.text_frame.margin_right = 0
    tp(pill1.text_frame, flow_badge, sz=8, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    P2.append((pill1.shape_id, 80, 'zm'))

    # ─── 4. LEFT COLUMN: INGESTION PIPELINE (X: 0.68, W: 2.00) ────────
    sx, sw = 0.68, 2.00
    ing = data.get("ingestion", {})

    # Card 1: Specialist (Y: 1.52, H: 0.70)
    c1 = ing.get("card1", {"title": "Accounting Specialist", "sub": "Process Lead", "emote": "spec"})
    spec = bx(slide, sx, 1.52, sw, 0.70, C["card_bg"], C["card_bd"], r=4000); shadow(spec, 8000, 4000, 6000)
    sg, se = emote_r(slide, get_emote(c1.get("emote", "spec")), sx, 1.52, sw, 0.70, C["accent_g"], 0.44)
    st_ = tx(slide, sx + 0.12, 1.66, sw - 0.64, 0.42)
    tp(st_.text_frame, c1.get("title", "Accounting Specialist"), sz=9.6, b=True, c=C["t1"])
    tp(st_.text_frame, c1.get("sub", "Lead"), sz=7.4, c=C["t3"], sb=2)
    P2 += [(spec.shape_id, 120, 'fade'), (sg.shape_id, 140, 'fade'), (se.shape_id, 150, 'ep'), (st_.shape_id, 170, 'fade')]

    a0 = ad(slide, sx + sw / 2, 2.22, 2.58, C["blue"], w=LW); P2.append((a0.shape_id, 200, 'wd'))

    # Card 2: Extract Report (Y: 2.58, H: 1.26)
    c2 = ing.get("card2", {"title": "Extract Open-Item Report", "sub": "Qlik Sense extract | Exclude matched items.", "emote": "qlik"})
    qlik = bx(slide, sx, 2.58, sw, 1.26, C["accent_l"], C["card_bd"], r=5000); shadow(qlik, 10000, 5000, 8000)
    qs = fl(slide, sx, 2.66, 0.045, 1.10, C["blue"])
    qg, qe = emote_r(slide, get_emote(c2.get("emote", "qlik")), sx, 2.58, sw, 1.26, C["accent_g"], 0.48)
    qt = tx(slide, sx + 0.12, 2.74, sw - 0.66, 0.94)
    tp(qt.text_frame, c2.get("title", "Extract Open-Item Report"), sz=10.0, b=True, c=C["t1"])
    tp(qt.text_frame, c2.get("sub", ""), sz=7.5, c=C["t3"], sb=3.5)
    P2 += [(qlik.shape_id, 240, 'fade'), (qs.shape_id, 255, 'fade'), (qg.shape_id, 265, 'fade'), (qe.shape_id, 275, 'ep'), (qt.shape_id, 295, 'fade')]

    a1 = ad(slide, sx + sw / 2, 3.84, 4.22, C["blue"], w=LW); P2.append((a1.shape_id, 330, 'wd'))

    # Card 3: Filter (Y: 4.22, H: 0.64)
    c3 = ing.get("card3", {"title": "Filter by Company Code", "sub": "Scope validation by entity code", "emote": "filter"})
    filt = bx(slide, sx, 4.22, sw, 0.64, C["card_bg"], C["card_bd"], r=4000); shadow(filt, 7000, 3500, 6000)
    fg, fe = emote_r(slide, get_emote(c3.get("emote", "filter")), sx, 4.22, sw, 0.64, C["accent_g"], 0.38)
    ft = tx(slide, sx + 0.12, 4.34, sw - 0.60, 0.40)
    tp(ft.text_frame, c3.get("title", "Filter by Company Code"), sz=8.8, b=True, c=C["blue"])
    tp(ft.text_frame, c3.get("sub", ""), sz=7.0, c=C["t3"], sb=2)
    P2 += [(filt.shape_id, 360, 'fade'), (fg.shape_id, 370, 'fade'), (fe.shape_id, 380, 'ep'), (ft.shape_id, 395, 'fade')]

    a2 = ad(slide, sx + sw / 2, 4.86, 5.24, C["blue"], w=LW); P2.append((a2.shape_id, 420, 'wd'))

    # Card 4: List / Repository (Y: 5.24, H: 1.82, ends at 7.06)
    c4 = ing.get("card4", {"title": "List of open items with classification", "sub": "Consolidated open delta ready for triage taxonomy.\nMaps unreconciled line items into operational resolution tracks.", "emote": "list"})
    lst = bx(slide, sx, 5.24, sw, 1.82, C["card_bg"], C["card_bd"], r=5000); shadow(lst, 12000, 6000, 8000)
    ls = fl(slide, sx, 5.32, 0.045, 1.66, C["blue"])
    lg, le = emote_r(slide, get_emote(c4.get("emote", "list")), sx, 5.24, sw, 1.82, C["accent_g"], 0.50)
    lt = tx(slide, sx + 0.12, 5.40, sw - 0.66, 1.50)
    tp(lt.text_frame, c4.get("title", "List of open items"), sz=10.0, b=True, c=C["t1"])
    tp(lt.text_frame, c4.get("sub", ""), sz=7.5, c=C["t3"], sb=3.5)
    P2 += [(lst.shape_id, 450, 'fade'), (ls.shape_id, 460, 'fade'), (lg.shape_id, 470, 'fade'), (le.shape_id, 480, 'ep'), (lt.shape_id, 500, 'fade')]

    # ─── 5. DECISION FORK (3-way spine from List to 3 Tracks) ──────────
    fx1 = sx + sw           # 2.68
    fxm = 2.88              # vertical spine
    fx2 = 3.08              # branch entry
    stem = ln(slide, fx1, 5.40, fxm, 5.40, C["purple"], w=LW)
    spine = ln(slide, fxm, 2.75, fxm, 6.54, C["purple"], w=LW)
    br1 = ar(slide, fxm, 2.75, fx2, C["purple"], w=LW)
    br2 = ar(slide, fxm, 5.00, fx2, C["purple"], w=LW)
    br3 = ar(slide, fxm, 6.54, fx2, C["purple"], w=LW)
    P3 += [(stem.shape_id, 0, 'wr'), (spine.shape_id, 80, 'fade'), (br1.shape_id, 120, 'wr'), (br2.shape_id, 140, 'wr'), (br3.shape_id, 160, 'wr')]

    # ─── 6. OPERATIONAL TRACKS ─────────────────────────────────────────
    tracks = data.get("tracks", {})
    t1_data = tracks.get("track1", {})
    t2_data = tracks.get("track2", {})
    t3_data = tracks.get("track3", {})

    # TRACK 1 (TOP): Posting not found (Y: 2.20, H: 1.10, center 2.75)
    tx1, tw1 = 3.08, 1.78
    pnf = bx(slide, tx1, 2.20, tw1, 1.10, C["card_bg"], C["card_bd"], r=5000); shadow(pnf, 8000, 4000, 6000)
    pnfs = fl(slide, tx1, 2.28, 0.045, 0.94, C["blue"])
    pnfg, pnfe = emote_r(slide, get_emote(t1_data.get("emote", "pnf")), tx1, 2.20, tw1, 1.10, C["accent_g"], 0.44)
    pnft = tx(slide, tx1 + 0.12, 2.38, tw1 - 0.62, 0.74)
    tp(pnft.text_frame, t1_data.get("title", "Posting not found"), sz=10.0, b=True, c=C["t1"])
    tp(pnft.text_frame, t1_data.get("sub", "Kernel out of scope line"), sz=7.5, c=C["t3"], sb=2.5)
    P3 += [(pnf.shape_id, 200, 'fade'), (pnfs.shape_id, 210, 'fade'), (pnfg.shape_id, 220, 'fade'), (pnfe.shape_id, 230, 'ep'), (pnft.shape_id, 250, 'fade')]

    # Symmetrical Sub-Fork to 1A and 1B
    sfx1, sfxm, sfx2 = tx1 + tw1, 5.02, 5.18
    sfs = ln(slide, sfx1, 2.75, sfxm, 2.75, C["purple"], w=LW)
    sfsp = ln(slide, sfxm, 2.03, sfxm, 3.47, C["purple"], w=LW)
    sfta = ar(slide, sfxm, 2.03, sfx2, C["purple"], w=LW)
    sftb = ar(slide, sfxm, 3.47, sfx2, C["purple"], w=LW)
    P4 += [(sfs.shape_id, 0, 'wr'), (sfsp.shape_id, 40, 'fade'), (sfta.shape_id, 70, 'wr'), (sftb.shape_id, 90, 'wr')]

    # Row 1A: IDoc / OCR (Y: 1.52, H: 1.02)
    iaw, hx, hw = 1.48, 6.92, 1.90
    t1a = t1_data.get("branch_a", {})
    ia = bx(slide, sfx2, 1.52, iaw, 1.02, C["card_bg"], C["card_bd"], r=4000)
    iag, iae = emote_r(slide, get_emote(t1a.get("cond_emote", "idoc")), sfx2, 1.52, iaw, 1.02, C["accent_g"], 0.42)
    iat = tx(slide, sfx2 + 0.10, 1.64, iaw - 0.54, 0.78)
    tp(iat.text_frame, t1a.get("cond_title", "IDoc / OCR Issue"), sz=9.2, b=True, c=C["t1"])
    tp(iat.text_frame, t1a.get("cond_sub", "Interface syntax failure"), sz=7.2, c=C["t3"], sb=2.5)

    ia_a = ar(slide, sfx2 + iaw, 2.03, hx, C["teal"], w=LW)

    hwi = bx(slide, hx, 1.52, hw, 1.02, C["card_bg"], C["card_bd"], r=4000)
    hwic = ci(slide, hx + 0.14, 2.03, 0.11, C["teal"])
    tp(hwic.text_frame, "›", sz=11, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    hwig, hwie = emote_r(slide, get_emote(t1a.get("act_emote", "hwi")), hx, 1.52, hw, 1.02, C["accent_g"], 0.42)
    hwit = tx(slide, hx + 0.30, 1.64, hw - 0.78, 0.78)
    tp(hwit.text_frame, t1a.get("act_title", "Troubleshoot IDoc Using HWI"), sz=9.0, b=True, c=C["t1"])
    tp(hwit.text_frame, t1a.get("act_sub", "Execute SAP Hand Work Instructions"), sz=7.2, c=C["t3"], sb=2.5)
    P4 += [(ia.shape_id, 120, 'fade'), (iag.shape_id, 130, 'fade'), (iae.shape_id, 140, 'ep'), (iat.shape_id, 150, 'fade'),
           (ia_a.shape_id, 170, 'wr'),
           (hwi.shape_id, 190, 'fade'), (hwic.shape_id, 200, 'zm'), (hwig.shape_id, 210, 'fade'), (hwie.shape_id, 220, 'ep'), (hwit.shape_id, 230, 'fade')]

    # Row 1B: No EDI / Non-SAP (Y: 2.96, H: 1.02)
    t1b = t1_data.get("branch_b", {})
    ib = bx(slide, sfx2, 2.96, iaw, 1.02, C["card_bg"], C["card_bd"], r=4000)
    ibg, ibe = emote_r(slide, get_emote(t1b.get("cond_emote", "no_edi")), sfx2, 2.96, iaw, 1.02, C["accent_g"], 0.42)
    ibt = tx(slide, sfx2 + 0.10, 3.08, iaw - 0.54, 0.78)
    tp(ibt.text_frame, t1b.get("cond_title", "No EDI / Non-SAP"), sz=9.2, b=True, c=C["t1"])
    tp(ibt.text_frame, t1b.get("cond_sub", "Paper drop / legacy format"), sz=7.2, c=C["t3"], sb=2.5)

    ib_a = ar(slide, sfx2 + iaw, 3.47, hx, C["blue"], w=LW)

    inv = bx(slide, hx, 2.96, hw, 1.02, C["card_bg"], C["card_bd"], r=4000)
    invc = ci(slide, hx + 0.14, 3.47, 0.11, C["blue"])
    tp(invc.text_frame, "›", sz=11, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    invg, inve = emote_r(slide, get_emote(t1b.get("act_emote", "inv")), hx, 2.96, hw, 1.02, C["accent_g"], 0.42)
    invt = tx(slide, hx + 0.30, 3.08, hw - 0.78, 0.78)
    tp(invt.text_frame, t1b.get("act_title", "Request / Retrieve Invoice Copy"), sz=9.0, b=True, c=C["t1"])
    tp(invt.text_frame, t1b.get("act_sub", "Auto-fetch PDF via OCR matching"), sz=7.2, c=C["t3"], sb=2.5)
    P4 += [(ib.shape_id, 260, 'fade'), (ibg.shape_id, 270, 'fade'), (ibe.shape_id, 280, 'ep'), (ibt.shape_id, 290, 'fade'),
           (ib_a.shape_id, 310, 'wr'),
           (inv.shape_id, 330, 'fade'), (invc.shape_id, 340, 'zm'), (invg.shape_id, 350, 'fade'), (inve.shape_id, 360, 'ep'), (invt.shape_id, 370, 'fade')]

    # Row 2: AP-AR Sign Issue (Y: 4.46, H: 1.08)
    t2x, t2w = 3.08, 1.94
    a2x, a2w = 5.38, 2.48
    t2 = bx(slide, t2x, 4.46, t2w, 1.08, C["card_bg"], C["card_bd"], r=5000); shadow(t2, 8000, 4000, 6000)
    t2s = fl(slide, t2x, 4.54, 0.045, 0.92, C["amber"])
    t2g, t2e = emote_r(slide, get_emote(t2_data.get("cond_emote", "ap_ar")), t2x, 4.46, t2w, 1.08, C["accent_g"], 0.44)
    t2t = tx(slide, t2x + 0.12, 4.64, t2w - 0.64, 0.72)
    tp(t2t.text_frame, t2_data.get("cond_title", "Investigate AP-AR sign issue"), sz=10.0, b=True, c=C["t1"])
    tp(t2t.text_frame, t2_data.get("cond_sub", "AR cleared, AP remains open (+/−)"), sz=7.5, c=C["t3"], sb=3)

    t2a = ar(slide, t2x + t2w, 5.00, a2x, C["amber"], w=LW)

    a2 = bx(slide, a2x, 4.46, a2w, 1.08, C["card_bg"], C["card_bd"], r=4000)
    a2c = ci(slide, a2x + 0.14, 5.00, 0.11, C["amber"])
    tp(a2c.text_frame, "›", sz=11, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    a2g, a2e = emote_r(slide, get_emote(t2_data.get("act_emote", "review")), a2x, 4.46, a2w, 1.08, C["accent_g"], 0.44)
    a2t = tx(slide, a2x + 0.30, 4.64, a2w - 0.80, 0.72)
    tp(a2t.text_frame, t2_data.get("act_title", "Review & Analyze Issue"), sz=9.6, b=True, c=C["t1"])
    tp(a2t.text_frame, t2_data.get("act_sub", "Investigate discrepancy; post clearing journal"), sz=7.4, c=C["t3"], sb=3)
    P4 += [(t2.shape_id, 400, 'fade'), (t2s.shape_id, 410, 'fade'), (t2g.shape_id, 420, 'fade'), (t2e.shape_id, 430, 'ep'), (t2t.shape_id, 440, 'fade'),
           (t2a.shape_id, 460, 'wr'),
           (a2.shape_id, 480, 'fade'), (a2c.shape_id, 490, 'zm'), (a2g.shape_id, 500, 'fade'), (a2e.shape_id, 510, 'ep'), (a2t.shape_id, 520, 'fade')]

    # Row 3: Cash Allocated (Y: 6.00, H: 1.08, ends at 7.08)
    t3 = bx(slide, t2x, 6.00, t2w, 1.08, C["card_bg"], C["card_bd"], r=5000); shadow(t3, 8000, 4000, 6000)
    t3s = fl(slide, t2x, 6.08, 0.045, 0.92, C["rose"])
    t3g, t3e = emote_r(slide, get_emote(t3_data.get("cond_emote", "cash")), t2x, 6.00, t2w, 1.08, C["accent_g"], 0.44)
    t3t = tx(slide, t2x + 0.12, 6.18, t2w - 0.64, 0.72)
    tp(t3t.text_frame, t3_data.get("cond_title", "Cash to allocated / AP paid"), sz=10.0, b=True, c=C["t1"])
    tp(t3t.text_frame, t3_data.get("cond_sub", "AR unapplied in reciprocal ERP"), sz=7.5, c=C["t3"], sb=3)

    t3a = ar(slide, t2x + t2w, 6.54, a2x, C["rose"], w=LW)

    a3 = bx(slide, a2x, 6.00, a2w, 1.08, C["card_bg"], C["card_bd"], r=4000)
    a3c = ci(slide, a2x + 0.14, 6.54, 0.11, C["rose"])
    tp(a3c.text_frame, "›", sz=11, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    a3g, a3e = emote_r(slide, get_emote(t3_data.get("act_emote", "waiting")), a2x, 6.00, a2w, 1.08, C["accent_g"], 0.44)
    a3t = tx(slide, a2x + 0.30, 6.18, a2w - 0.80, 0.72)
    tp(a3t.text_frame, t3_data.get("act_title", "Waiting for Counterparty Action"), sz=9.6, b=True, c=C["t1"])
    tp(a3t.text_frame, t3_data.get("act_sub", "Pending reciprocal entity ledger clearing"), sz=7.4, c=C["t3"], sb=3)
    P4 += [(t3.shape_id, 550, 'fade'), (t3s.shape_id, 560, 'fade'), (t3g.shape_id, 570, 'fade'), (t3e.shape_id, 580, 'ep'), (t3t.shape_id, 590, 'fade'),
           (t3a.shape_id, 610, 'wr'),
           (a3.shape_id, 630, 'fade'), (a3c.shape_id, 640, 'zm'), (a3g.shape_id, 650, 'fade'), (a3e.shape_id, 660, 'ep'), (a3t.shape_id, 670, 'fade')]

    # ─── 7. CONVERGENCE TO GOVERNANCE ──────────────────────────────────
    ext2 = ln(slide, a2x + a2w, 5.00, 8.82, 5.00, C["purple"], w=LW)
    ext3 = ln(slide, a2x + a2w, 6.54, 8.82, 6.54, C["purple"], w=LW)
    cl1 = ln(slide, 8.82, 2.03, 9.02, 2.03, C["purple"], w=LW)
    cl2 = ln(slide, 8.82, 3.47, 9.02, 3.47, C["purple"], w=LW)
    cl3 = ln(slide, 8.82, 5.00, 9.02, 5.00, C["purple"], w=LW)
    cl4 = ln(slide, 8.82, 6.54, 9.02, 6.54, C["purple"], w=LW)
    csp = ln(slide, 9.02, 2.03, 9.02, 6.54, C["purple"], w=LW)
    ca = ar(slide, 9.02, 2.12, 9.28, C["purple"], w=LW)
    P5 += [(ext2.shape_id, 0, 'wr'), (ext3.shape_id, 20, 'wr'),
           (cl1.shape_id, 40, 'wr'), (cl2.shape_id, 60, 'wr'), (cl3.shape_id, 80, 'wr'), (cl4.shape_id, 100, 'wr'),
           (csp.shape_id, 120, 'fade'), (ca.shape_id, 160, 'wr')]

    # ─── 8. GOVERNANCE COLUMN (X: 9.28, W: 3.45) ──────────────────────
    gx, gw = 9.28, 3.45
    gov = data.get("governance", {})

    # Card 1: Gap Detected (Y: 1.52, H: 1.20)
    g1 = gov.get("card1", {"tag": "GAP DETECTED", "title": "Reconciliation Discrepancy", "sub": "Open-item delta isolated between reciprocal entities.", "emote": "gap"})
    disc = bx(slide, gx, 1.52, gw, 1.20, C["accent_l"], C["card_bd"], r=5000); shadow(disc, 10000, 5000, 8000)
    dt = bx(slide, gx + 0.14, 1.64, 0.96, 0.22, C["rose"], None, r=3000)
    tp(dt.text_frame, g1.get("tag", "GAP DETECTED"), sz=7.0, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    dg, de = emote_r(slide, get_emote(g1.get("emote", "gap")), gx, 1.52, gw, 1.20, C["accent_g"], 0.50)
    dtx = tx(slide, gx + 0.14, 1.96, gw - 0.76, 0.70)
    tp(dtx.text_frame, g1.get("title", "Reconciliation Discrepancy"), sz=10.8, b=True, c=C["t1"])
    tp(dtx.text_frame, g1.get("sub", ""), sz=7.8, c=C["t3"], sb=3)
    P5 += [(disc.shape_id, 200, 'fade'), (dt.shape_id, 220, 'fade'), (dg.shape_id, 230, 'fade'), (de.shape_id, 240, 'ep'), (dtx.shape_id, 260, 'fade')]

    pa1 = ad(slide, gx + gw / 2, 2.72, 3.16, C["rose"], w=LW); P5.append((pa1.shape_id, 310, 'wd'))

    # Card 2: Action Notification (Y: 3.16, H: 1.20)
    g2 = gov.get("card2", {"tag": "ACTION NOTIFICATION", "title": "Send Action Notification (Email)", "sub": "Structured notice sent to counterparty accounting lead.", "emote": "notif"})
    not_ = bx(slide, gx, 3.16, gw, 1.20, C["card_bg"], C["card_bd"], r=5000); shadow(not_, 10000, 5000, 8000)
    nt = bx(slide, gx + 0.14, 3.28, 1.40, 0.22, C["rose"], None, r=3000)
    tp(nt.text_frame, g2.get("tag", "ACTION NOTIFICATION"), sz=7.0, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    ng, ne = emote_r(slide, get_emote(g2.get("emote", "notif")), gx, 3.16, gw, 1.20, C["accent_g"], 0.50)
    ntx = tx(slide, gx + 0.14, 3.60, gw - 0.76, 0.70)
    tp(ntx.text_frame, g2.get("title", "Send Action Notification"), sz=10.8, b=True, c=C["t1"])
    tp(ntx.text_frame, g2.get("sub", ""), sz=7.8, c=C["t3"], sb=3)
    P5 += [(not_.shape_id, 360, 'fade'), (nt.shape_id, 380, 'fade'), (ng.shape_id, 390, 'fade'), (ne.shape_id, 400, 'ep'), (ntx.shape_id, 420, 'fade')]

    pa2 = ad(slide, gx + gw / 2, 4.36, 4.80, C["red"], w=LW); P5.append((pa2.shape_id, 470, 'wd'))

    # Card 3: Escalation Matrix (Y: 4.80, H: 2.28, ends at 7.08)
    g3 = gov.get("card3", {"tag": "MULTI-TIER ESCALATION", "title": "No Response → Initiate Escalation (Matrix)", "emote": "esc"})
    esc = bx(slide, gx, 4.80, gw, 2.28, C["accent_l"], C["card_bd"], r=5000); shadow(esc, 14000, 7000, 9000)
    et = bx(slide, gx + 0.14, 4.94, 1.44, 0.22, C["red"], None, r=3000)
    tp(et.text_frame, g3.get("tag", "MULTI-TIER ESCALATION"), sz=7.0, b=True, c=C["card_bg"], a=PP_ALIGN.CENTER)
    eg, ee = emote_r(slide, get_emote(g3.get("emote", "esc")), gx, 4.80, gw, 2.28, C["accent_g"], 0.52)
    etx = tx(slide, gx + 0.14, 5.26, gw - 0.76, 1.74)
    tp(etx.text_frame, g3.get("title", "No Response → Initiate Escalation"), sz=10.2, b=True, c=C["red"])

    tiers = g3.get("tiers", [
        ("L1", "48h Inaction", "Accounting Lead / Processor", "Initial SLA alert; re-verify unmatched ledger delta."),
        ("L2", "96h Inaction", "Shared Services Manager", "Shared services escalation; bilateral review call."),
        ("L3", ">5d / Close", "Entity Finance Director", "Executive sign-off; post un-cleared accrual & audit note.")
    ])
    for tier, hrs, own, detail in tiers:
        p = etx.text_frame.add_paragraph()
        p.space_before = Pt(5.0); p.line_spacing = Pt(11.5)
        r1 = p.add_run(); r1.text = tier + " "; r1.font.size = Pt(7.8); r1.font.bold = True; r1.font.color.rgb = C["red"]
        r2 = p.add_run(); r2.text = f"({hrs}): "; r2.font.size = Pt(7.8); r2.font.bold = True; r2.font.color.rgb = C["t1"]
        r3 = p.add_run(); r3.text = own + "\n"; r3.font.size = Pt(7.8); r3.font.bold = True; r3.font.color.rgb = C["t2"]
        r4 = p.add_run(); r4.text = "    " + detail; r4.font.size = Pt(7.0); r4.font.color.rgb = C["t3"]
    P5 += [(esc.shape_id, 520, 'fade'), (et.shape_id, 540, 'fade'), (eg.shape_id, 550, 'fade'), (ee.shape_id, 560, 'ep'), (etx.shape_id, 580, 'fade')]

    # ─── 9. COMPILE & SAVE ─────────────────────────────────────────────
    eng.add(P1); eng.add(P2); eng.add(P3); eng.add(P4); eng.add(P5)
    eng.build(slide)
    prs.save(out_path)
    return out_path
