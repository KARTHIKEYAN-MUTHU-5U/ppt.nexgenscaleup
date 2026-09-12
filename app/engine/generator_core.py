#!/usr/bin/env python3
"""
Executive PowerPoint Presentation Engine - Core Utilities
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Provides OpenXML DrawingML 2.0 pt connectors, native PresentationML
directional wipe animations, halo glows, and shape styling.
"""

import os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

LW = 2.0  # Standardized 2.0 pt line width everywhere

# ═══ COLOR PALETTES ════════════════════════════════════════════════════
PALETTES = {
    "navy_corporate": {
        "name": "Navy Corporate (Philips Style)",
        "hdr_bg": RGBColor(0x13, 0x23, 0x3D),
        "canvas_bg": RGBColor(0xF4, 0xF6, 0xF9),
        "card_bg": RGBColor(0xFF, 0xFF, 0xFF),
        "card_bd": RGBColor(0xD8, 0xE2, 0xEC),
        "t1": RGBColor(0x10, 0x1D, 0x30),
        "t2": RGBColor(0x2D, 0x3F, 0x55),
        "t3": RGBColor(0x5E, 0x77, 0x94),
        "t4": RGBColor(0x9A, 0xB0, 0xC6),
        "accent": RGBColor(0x0B, 0x5C, 0xAD),
        "accent_l": RGBColor(0xEE, 0xF5, 0xFC),
        "accent_g": RGBColor(0xD6, 0xEA, 0xFA),
        "stripe": RGBColor(0x00, 0xA8, 0x84),
        "blue": RGBColor(0x0B, 0x5C, 0xAD),
        "amber": RGBColor(0xD9, 0x77, 0x06),
        "teal": RGBColor(0x00, 0x8C, 0x95),
        "rose": RGBColor(0xD9, 0x46, 0x8D),
        "red": RGBColor(0xDC, 0x26, 0x26),
        "purple": RGBColor(0x63, 0x66, 0xF1),
        "ctn_bd": RGBColor(0xBF, 0xE3, 0xE6),
    },
    "obsidian_sapphire": {
        "name": "Obsidian & Sapphire (Modern Tech)",
        "hdr_bg": RGBColor(0x0E, 0x17, 0x26),
        "canvas_bg": RGBColor(0xF1, 0xF5, 0xF9),
        "card_bg": RGBColor(0xFF, 0xFF, 0xFF),
        "card_bd": RGBColor(0xCB, 0xD5, 0xE1),
        "t1": RGBColor(0x0F, 0x17, 0x2A),
        "t2": RGBColor(0x33, 0x41, 0x55),
        "t3": RGBColor(0x64, 0x74, 0x8B),
        "t4": RGBColor(0x94, 0xA3, 0xB8),
        "accent": RGBColor(0x02, 0x84, 0xC7),
        "accent_l": RGBColor(0xF0, 0xF9, 0xFF),
        "accent_g": RGBColor(0xBA, 0xE6, 0xFD),
        "stripe": RGBColor(0x38, 0xBD, 0xF8),
        "blue": RGBColor(0x02, 0x84, 0xC7),
        "amber": RGBColor(0xF5, 0x9E, 0x0B),
        "teal": RGBColor(0x0D, 0x94, 0x88),
        "rose": RGBColor(0xE1, 0x1D, 0x48),
        "red": RGBColor(0xEF, 0x44, 0x44),
        "purple": RGBColor(0x8B, 0x5C, 0xF6),
        "ctn_bd": RGBColor(0x93, 0xC5, 0xFD),
    },
    "slate_emerald": {
        "name": "Slate & Emerald (Sustainability & Strategy)",
        "hdr_bg": RGBColor(0x0F, 0x29, 0x1E),
        "canvas_bg": RGBColor(0xF5, 0xF7, 0xF6),
        "card_bg": RGBColor(0xFF, 0xFF, 0xFF),
        "card_bd": RGBColor(0xD1, 0xDC, 0xD6),
        "t1": RGBColor(0x0D, 0x1F, 0x17),
        "t2": RGBColor(0x27, 0x3E, 0x33),
        "t3": RGBColor(0x57, 0x70, 0x64),
        "t4": RGBColor(0x8D, 0xA8, 0x9C),
        "accent": RGBColor(0x05, 0x96, 0x69),
        "accent_l": RGBColor(0xEC, 0xFD, 0xF5),
        "accent_g": RGBColor(0xA7, 0xF3, 0xD0),
        "stripe": RGBColor(0x10, 0xB9, 0x81),
        "blue": RGBColor(0x02, 0x84, 0xC7),
        "amber": RGBColor(0xD9, 0x77, 0x06),
        "teal": RGBColor(0x05, 0x96, 0x69),
        "rose": RGBColor(0xBE, 0x18, 0x5D),
        "red": RGBColor(0xDC, 0x26, 0x26),
        "purple": RGBColor(0x63, 0x66, 0xF1),
        "ctn_bd": RGBColor(0xA7, 0xF3, 0xD0),
    }
}

# ═══ DRAWINGML & SHAPE UTILITIES ══════════════════════════════════════
def fs(s, c): 
    s.fill.solid()
    s.fill.fore_color.rgb = c

def sb(s, c, w=0.75): 
    s.line.color.rgb = c
    s.line.width = Pt(w)

def nb(s): 
    s.line.fill.background()

def sc(s, r=10000):
    """Set shape corner radius adj."""
    pg = s._element.find('.//' + qn('a:prstGeom'))
    if pg is None: return
    av = pg.find(qn('a:avLst'))
    if av is None: av = etree.SubElement(pg, qn('a:avLst'))
    else:
        for ch in list(av): av.remove(ch)
    gd = etree.SubElement(av, qn('a:gd'))
    gd.set('name', 'adj')
    gd.set('fmla', f'val {r}')

def shadow(s, bl=20000, d=10000, a=12000):
    """Add subtle executive drop shadow."""
    spPr = s._element.find(qn('p:spPr'))
    if spPr is None: spPr = s._element.find(qn('a:spPr'))
    if spPr is None: return
    el = spPr.find(qn('a:effectLst'))
    if el is None: el = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(el, qn('a:outerShdw'))
    sh.set('blurRad', str(bl))
    sh.set('dist', str(d))
    sh.set('dir', '5400000')
    sh.set('rotWithShape', '0')
    clr = etree.SubElement(sh, qn('a:srgbClr'))
    clr.set('val', '000000')
    alp = etree.SubElement(clr, qn('a:alpha'))
    alp.set('val', str(a))

def dash_bd(s, c, w=1.5, d='lgDash'):
    """Add dashed stroke border."""
    s.line.color.rgb = c
    s.line.width = Pt(w)
    spPr = s._element.find(qn('p:spPr'))
    ln = spPr.find(qn('a:ln'))
    if ln is None: ln = etree.SubElement(spPr, qn('a:ln'))
    pd = ln.find(qn('a:prstDash'))
    if pd is None: pd = etree.SubElement(ln, qn('a:prstDash'))
    pd.set('val', d)

def bx(sl, x, y, w, h, bg, bd=None, bw=0.75, r=10000):
    """Add styled rounded rectangle card."""
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fs(s, bg)
    if bd: sb(s, bd, bw)
    else: nb(s)
    sc(s, r)
    s.text_frame.word_wrap = True
    s.text_frame.auto_size = None
    s.text_frame.margin_left = Inches(0.12)
    s.text_frame.margin_right = Inches(0.06)
    s.text_frame.margin_top = Inches(0.08)
    s.text_frame.margin_bottom = Inches(0.08)
    return s

def fl(sl, x, y, w, h, bg):
    """Add flat rectangle / accent block."""
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fs(s, bg)
    nb(s)
    return s

def ci(sl, cx, cy, r, bg):
    """Add centered circle."""
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(r * 2), Inches(r * 2))
    fs(s, bg)
    nb(s)
    s.text_frame.margin_left = s.text_frame.margin_right = 0
    s.text_frame.margin_top = s.text_frame.margin_bottom = 0
    return s

def pk(sl, p, x, y, w, h):
    """Add picture image / GIF."""
    return sl.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))

def tx(sl, x, y, w, h):
    """Add empty textbox."""
    t = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.word_wrap = True
    t.text_frame.margin_left = t.text_frame.margin_right = 0
    t.text_frame.margin_top = t.text_frame.margin_bottom = 0
    return t

def tp(tf, text, sz=9, b=False, c=None, f='Aptos', a=PP_ALIGN.LEFT, sb=0):
    """Add paragraph with executive typography."""
    p = tf.paragraphs[0] if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == '' else tf.add_paragraph()
    p.alignment = a
    if sb: p.space_before = Pt(sb)
    p.line_spacing = Pt(sz * 1.22)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = b
    r.font.name = f
    if c: r.font.color.rgb = c
    return p

# ═══ DRAWINGML 2.0 PT CONNECTORS & ARROWS ═════════════════════════════
def conn_line(sl, x1, y1, x2, y2, c, w=LW):
    """Draw a straight connector line with NO arrow, exact width w pt."""
    conn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = c
    conn.line.width = Pt(w)
    return conn

def conn_arrow(sl, x1, y1, x2, y2, c, w=LW, head='triangle', head_w='med', head_len='med'):
    """Draw a straight connector arrow from (x1, y1) to (x2, y2) with exact width w pt and uniform arrowhead."""
    conn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = c
    conn.line.width = Pt(w)
    spPr = conn._element.spPr
    ln = spPr.find(qn('a:ln'))
    if ln is None: ln = etree.SubElement(spPr, qn('a:ln'))
    te = etree.SubElement(ln, qn('a:tailEnd'))
    te.set('type', head)
    te.set('w', head_w)
    te.set('len', head_len)
    return conn

def ad(sl, cx, y1, y2, c, w=LW):
    """Uniform Downward Arrow."""
    return conn_arrow(sl, cx, y1, cx, y2, c, w=w)

def ar(sl, x1, cy, x2, c, w=LW):
    """Uniform Rightward Arrow."""
    return conn_arrow(sl, x1, cy, x2, cy, c, w=w)

def ln(sl, x1, y1, x2, y2, c, w=LW):
    """Uniform Line Segment."""
    return conn_line(sl, x1, y1, x2, y2, c, w=w)

def emote_r(sl, img_path, card_x, card_y, card_w, card_h, glow_c, sz=0.46):
    """Place emote cleanly inside card on right side, with a soft glowing halo."""
    ex = card_x + card_w - sz - 0.10
    ey = card_y + (card_h - sz) / 2
    g = ci(sl, ex + sz / 2, ey + sz / 2, sz * 0.54, glow_c)
    shadow(g, bl=5000, d=0, a=8000)
    e = pk(sl, img_path, ex, ey, sz, sz)
    return g, e

# ═══ OPENXML ANIMATION ENGINE ═════════════════════════════════════════
class AnimationEngine:
    """Compiles PresentationML <p:timing> sequences."""
    NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'

    def __init__(self):
        self.sid = 0
        self.ph = []
        self.sp = []

    def _id(self):
        self.sid += 1
        return str(self.sid)

    def tg(self, t):
        return f'{{{self.NS}}}{t}'

    def add(self, items):
        self.ph.append(items)
        for s, _, _ in items:
            if s not in self.sp:
                self.sp.append(s)

    def build(self, slide):
        el = slide._element
        old = el.find(qn('p:timing'))
        if old is not None:
            el.remove(old)
        el.append(self._tree())

    def _tree(self):
        tm = etree.Element(self.tg('timing'))
        tl = etree.SubElement(tm, self.tg('tnLst'))
        rp = etree.SubElement(tl, self.tg('par'))
        rc = etree.SubElement(rp, self.tg('cTn'))
        rc.set('id', self._id()); rc.set('dur', 'indefinite'); rc.set('restart', 'never'); rc.set('nodeType', 'tmRoot')
        rh = etree.SubElement(rc, self.tg('childTnLst'))
        sq = etree.SubElement(rh, self.tg('seq'))
        sq.set('concurrent', '1'); sq.set('nextAc', 'seek')
        sc = etree.SubElement(sq, self.tg('cTn'))
        sc.set('id', self._id()); sc.set('dur', 'indefinite'); sc.set('nodeType', 'mainSeq')
        sh = etree.SubElement(sc, self.tg('childTnLst'))
        for p in self.ph:
            sh.append(self._ph(p))
        for ev, tg in [('onPrev', 'prevCondLst'), ('onNext', 'nextCondLst')]:
            cl = etree.SubElement(sq, self.tg(tg))
            cd = etree.SubElement(cl, self.tg('cond'))
            cd.set('evt', ev); cd.set('delay', '0')
            te = etree.SubElement(cd, self.tg('tgtEl'))
            etree.SubElement(te, self.tg('sldTgt'))
        bl = etree.SubElement(tm, self.tg('bldLst'))
        for s in self.sp:
            bp = etree.SubElement(bl, self.tg('bldP'))
            bp.set('spid', str(s)); bp.set('grpId', '0')
        return tm

    def _ph(self, items):
        gp = etree.Element(self.tg('par'))
        gc = etree.SubElement(gp, self.tg('cTn'))
        gc.set('id', self._id()); gc.set('fill', 'hold')
        st = etree.SubElement(gc, self.tg('stCondLst'))
        cd = etree.SubElement(st, self.tg('cond'))
        cd.set('delay', '0')
        ch = etree.SubElement(gc, self.tg('childTnLst'))
        for i, (s, d, e) in enumerate(items):
            ch.append(self._an(s, d, e, 'clickEffect' if i == 0 else 'withEffect'))
        return gp

    def _an(self, spid, delay, eff, nt):
        pid, ps, flt, dur = '10', '0', 'fade', '850'
        if eff in ('wr', 'wipe_right'):
            pid, ps, flt, dur = '22', '8', 'wipe(left)', '750'
        elif eff in ('wd', 'wipe_down'):
            pid, ps, flt, dur = '22', '1', 'wipe(up)', '750'
        elif eff == 'zm':
            pid, ps, flt, dur = '53', '0', 'fade', '750'
        elif eff == 'ep':
            pid, ps, flt, dur = '53', '0', 'fade', '650'

        op = etree.Element(self.tg('par'))
        oc = etree.SubElement(op, self.tg('cTn'))
        oc.set('id', self._id()); oc.set('fill', 'hold')
        os = etree.SubElement(oc, self.tg('stCondLst'))
        ocd = etree.SubElement(os, self.tg('cond'))
        ocd.set('delay', str(delay))
        ich = etree.SubElement(oc, self.tg('childTnLst'))
        ap = etree.SubElement(ich, self.tg('par'))
        ac = etree.SubElement(ap, self.tg('cTn'))
        ac.set('id', self._id()); ac.set('presetID', pid); ac.set('presetClass', 'entr'); ac.set('presetSubtype', ps)
        ac.set('fill', 'hold'); ac.set('grpId', '0'); ac.set('nodeType', nt)
        ast = etree.SubElement(ac, self.tg('stCondLst'))
        acd = etree.SubElement(ast, self.tg('cond'))
        acd.set('delay', '0')
        ach = etree.SubElement(ac, self.tg('childTnLst'))
        se = etree.SubElement(ach, self.tg('set'))
        cb1 = etree.SubElement(se, self.tg('cBhvr'))
        ct1 = etree.SubElement(cb1, self.tg('cTn'))
        ct1.set('id', self._id()); ct1.set('dur', '1'); ct1.set('fill', 'hold')
        s1 = etree.SubElement(ct1, self.tg('stCondLst'))
        c1 = etree.SubElement(s1, self.tg('cond'))
        c1.set('delay', '0')
        tg1 = etree.SubElement(cb1, self.tg('tgtEl'))
        sp1 = etree.SubElement(tg1, self.tg('spTgt'))
        sp1.set('spid', str(spid))
        al = etree.SubElement(cb1, self.tg('attrNameLst'))
        an = etree.SubElement(al, self.tg('attrName'))
        an.text = 'style.visibility'
        to = etree.SubElement(se, self.tg('to'))
        sv = etree.SubElement(to, self.tg('strVal'))
        sv.set('val', 'visible')
        ae = etree.SubElement(ach, self.tg('animEffect'))
        ae.set('transition', 'in'); ae.set('filter', flt)
        cb2 = etree.SubElement(ae, self.tg('cBhvr'))
        ct2 = etree.SubElement(cb2, self.tg('cTn'))
        ct2.set('id', self._id()); ct2.set('dur', dur)
        tg2 = etree.SubElement(cb2, self.tg('tgtEl'))
        sp2 = etree.SubElement(tg2, self.tg('spTgt'))
        sp2.set('spid', str(spid))
        return op
